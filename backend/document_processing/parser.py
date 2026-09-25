import hashlib
import io
import re
from dataclasses import dataclass, field

from pypdf import PdfReader

try:
    import docx  # python-docx
except Exception:
    docx = None

ALLOWED_TYPES = {"pdf", "docx", "txt", "md", "csv"}
MAX_SIZE_MB = 25


@dataclass
class ParsedSection:
    section_id: str
    heading: str
    content: str
    page_ref: str = ""


@dataclass
class ParsedDocument:
    title: str
    file_type: str
    content_hash: str
    full_text: str
    sections: list[ParsedSection] = field(default_factory=list)
    page_count: int = 0
    warnings: list[str] = field(default_factory=list)


class DocumentValidationError(Exception):
    pass


def validate_file(filename: str, size_bytes: int, content: bytes, max_mb: int = MAX_SIZE_MB) -> str:
    if not filename or "." not in filename:
        raise DocumentValidationError("Filename must have an extension")
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in ALLOWED_TYPES:
        raise DocumentValidationError(f"Unsupported file type .{ext}. Allowed: {sorted(ALLOWED_TYPES)}")
    if size_bytes == 0 or len(content) == 0:
        raise DocumentValidationError("Empty document")
    if size_bytes > max_mb * 1024 * 1024:
        raise DocumentValidationError(f"File exceeds {max_mb}MB limit")
    if not content.strip():
        raise DocumentValidationError("Document has no extractable content")
    return ext


def content_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _split_sections(text: str, page_prefix: str = "") -> list[ParsedSection]:
    lines = text.splitlines()
    sections: list[ParsedSection] = []
    current_heading = "Introduction"
    current_buf: list[str] = []
    idx = 1

    heading_re = re.compile(
        r"^(?:section\s+[\d.]+|chapter\s+\d+|\d+(\.\d+)*\s+\S|#{1,3}\s+\S|[A-Z][A-Z\s\-]{5,})$",
        re.IGNORECASE,
    )

    def flush():
        nonlocal idx, current_buf, current_heading
        body = "\n".join(current_buf).strip()
        if body:
            sections.append(
                ParsedSection(
                    section_id=f"{page_prefix}{idx}",
                    heading=current_heading.strip(),
                    content=body,
                    page_ref=page_prefix or "",
                )
            )
            idx += 1
        current_buf = []

    for line in lines:
        if heading_re.match(line.strip()) and len(line.strip()) < 120:
            flush()
            current_heading = line.strip()
        else:
            current_buf.append(line)
    flush()

    if not sections and text.strip():
        sections.append(ParsedSection(section_id="1", heading="Full Document", content=text.strip()))
    return sections


def parse_pdf(data: bytes) -> ParsedDocument:
    reader = PdfReader(io.BytesIO(data))
    page_texts = []
    for i, page in enumerate(reader.pages, start=1):
        t = page.extract_text() or ""
        if t.strip():
            page_texts.append(f"[Page {i}]\n{t}")
    full = "\n\n".join(page_texts).strip()
    if not full:
        raise DocumentValidationError("PDF has no extractable text (possibly scanned)")
    sections = _split_sections(full, page_prefix="P")
    return ParsedDocument(
        title=(reader.metadata.title if reader.metadata and reader.metadata.title else "Untitled PDF"),
        file_type="pdf",
        content_hash=content_hash(data),
        full_text=full,
        sections=sections,
        page_count=len(reader.pages),
    )


def parse_docx(data: bytes) -> ParsedDocument:
    if docx is None:
        raise DocumentValidationError("python-docx not installed")
    document = docx.Document(io.BytesIO(data))
    parts = []
    for p in document.paragraphs:
        style = (p.style.name or "") if p.style else ""
        text = p.text.rstrip()
        if not text:
            continue
        if style.startswith("Heading") or style.startswith("Title"):
            parts.append(f"\n## {text}")
        else:
            parts.append(text)
    full = "\n".join(parts).strip()
    if not full:
        raise DocumentValidationError("DOCX has no extractable text")
    sections = _split_sections(full)
    return ParsedDocument(
        title=document.core_properties.title or "Untitled DOCX",
        file_type="docx",
        content_hash=content_hash(data),
        full_text=full,
        sections=sections,
        page_count=len(document.paragraphs),
    )


def parse_text(data: bytes) -> ParsedDocument:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1", errors="replace")
    text = text.strip()
    if not text:
        raise DocumentValidationError("Text file is empty")
    return ParsedDocument(
        title="Untitled Text",
        file_type="txt",
        content_hash=content_hash(data),
        full_text=text,
        sections=_split_sections(text),
        page_count=1,
    )


def parse_document(filename: str, data: bytes, max_mb: int = MAX_SIZE_MB) -> ParsedDocument:
    ext = validate_file(filename, len(data), data, max_mb=max_mb)
    if ext == "pdf":
        return parse_pdf(data)
    if ext == "docx":
        return parse_docx(data)
    return parse_text(data)
